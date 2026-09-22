from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Inches


class PPTGenerator:
    def __init__(self, output_dir: str = './storage/artifacts') -> None:
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate(self, title: str, slides: list[dict[str, Any]] | None = None, **kwargs: Any) -> dict[str, Any]:
        presentation = Presentation()
        presentation.slide_width = 9144000
        presentation.slide_height = 5143500

        for index, slide_data in enumerate(slides or []):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            title_text = slide_data.get('title', f'Slide {index + 1}')
            title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(8.5), Inches(0.7))
            title_box.text_frame.text = title_text

            body_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.3), Inches(8.2), Inches(4.4))
            body_frame = body_box.text_frame
            for bullet_index, bullet in enumerate(slide_data.get('bullets', [])):
                paragraph = body_frame.paragraphs[0] if bullet_index == 0 else body_frame.add_paragraph()
                paragraph.text = bullet

        title_lower = (title or 'generated_deck').lower()
        if 'proposal' in title_lower:
            stem = 'generated_proposal'
        elif 'deck' in title_lower or 'presentation' in title_lower:
            stem = 'generated_deck'
        else:
            stem = 'generated_presentation'

        file_path = self.output_dir / f"{stem}.pptx"
        presentation.save(str(file_path))
        return {
            'status': 'completed',
            'artifact_type': 'pptx',
            'file_path': str(file_path),
            'title': title,
        }
