from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document


class DocumentGenerator:
    def __init__(self, output_dir: str = './storage/artifacts') -> None:
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate(self, title: str, sections: list[str] | None = None, body: str = '', **kwargs: Any) -> dict[str, Any]:
        doc = Document()
        doc.add_heading(title or 'Generated Document', level=0)
        for section in sections or []:
            doc.add_heading(section, level=1)
        if body:
            doc.add_paragraph(body)

        title_lower = (title or 'generated_document').lower()
        if 'proposal' in title_lower:
            stem = 'generated_proposal'
        elif 'deck' in title_lower or 'presentation' in title_lower:
            stem = 'generated_deck'
        else:
            stem = 'generated_document'

        file_path = self.output_dir / f"{stem}.docx"
        doc.save(str(file_path))
        return {
            'status': 'completed',
            'artifact_type': 'docx',
            'file_path': str(file_path),
            'title': title,
        }
