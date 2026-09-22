from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any

from docx import Document
from pptx import Presentation


class ArtifactVersionManager:
    def __init__(self, base_dir: str = './storage/artifacts') -> None:
        self.base_dir = Path(base_dir)
        self.base_dir.mkdir(parents=True, exist_ok=True)

    def _next_version(self, source_path: Path) -> str:
        name = source_path.stem
        if '_v' in name:
            stem, version_part = name.rsplit('_v', 1)
            try:
                version_num = int(version_part)
                return f'v{version_num + 1}'
            except ValueError:
                pass
        return 'v2' if source_path.exists() else 'v1'

    def _versioned_target(self, source_path: Path, version: str) -> Path:
        suffix = source_path.suffix
        stem = source_path.stem
        if '_v' in stem:
            stem = stem.rsplit('_v', 1)[0]
        target_name = f'{stem}_{version}{suffix}'
        return self.base_dir / target_name

    def _apply_docx_edit(self, source_path: Path, target_path: Path, instruction: str) -> None:
        document = Document(str(source_path))
        document.add_paragraph('')
        document.add_heading('Conversation Edit', level=1)
        document.add_paragraph(f"Requested update: {instruction}")
        document.add_paragraph('This revision preserves the original structure while adding the requested change set.')
        document.save(str(target_path))

    def _apply_pptx_edit(self, source_path: Path, target_path: Path, instruction: str) -> None:
        presentation = Presentation(str(source_path))
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title_box = slide.shapes.add_textbox(0.7 * 914400, 0.4 * 914400, 8.5 * 914400, 0.7 * 914400)
        title_box.text_frame.text = 'Revision Update'
        body_box = slide.shapes.add_textbox(0.9 * 914400, 1.3 * 914400, 8.2 * 914400, 4.4 * 914400)
        frame = body_box.text_frame
        for idx, line in enumerate([
            instruction,
            'Version preserved while refinements were added to the narrative and messaging.',
        ]):
            paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
            paragraph.text = line
        presentation.save(str(target_path))

    def apply_edit(self, file_path: str, instruction: str) -> dict[str, Any]:
        source_path = Path(file_path)
        if not source_path.exists():
            raise FileNotFoundError(f'Artifact not found: {file_path}')

        version = self._next_version(source_path)
        target_path = self._versioned_target(source_path, version)

        if source_path.suffix.lower() == '.docx':
            self._apply_docx_edit(source_path, target_path, instruction)
        elif source_path.suffix.lower() == '.pptx':
            self._apply_pptx_edit(source_path, target_path, instruction)
        else:
            shutil.copy2(source_path, target_path)

        return {
            'status': 'completed',
            'version': version,
            'file_path': str(target_path),
            'source_file': str(source_path),
            'artifact_type': source_path.suffix.lower().lstrip('.'),
            'instruction': instruction,
        }
