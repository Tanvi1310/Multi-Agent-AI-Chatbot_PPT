from __future__ import annotations

from collections import Counter
from pathlib import Path

from pptx import Presentation


class PPTAnalyzer:
    def analyze(self, file_path: str) -> dict:
        prs = Presentation(file_path)
        slides = []
        slide_text = []
        fonts = Counter()
        colors = Counter()

        for index, slide in enumerate(prs.slides, start=1):
            slide_data = {"slide_number": index, "texts": []}
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    text = shape.text.strip()
                    slide_data["texts"].append(text)
                    slide_text.append(text)
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if run.font.name:
                                fonts[run.font.name] += 1
                            if run.font.color and getattr(run.font.color, "rgb", None):
                                colors[f"#{run.font.color.rgb}"] += 1
            slides.append(slide_data)

        return {
            "source_type": "pptx",
            "file_name": Path(file_path).name,
            "title": f"{Path(file_path).stem} ({len(slides)} slides)",
            "text": "\n".join(slide_text),
            "slide_count": len(slides),
            "slides": slides,
            "style": {
                "fonts": dict(fonts.most_common(5)),
                "colors": dict(colors.most_common(5)),
            },
        }
