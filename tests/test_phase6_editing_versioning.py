from pathlib import Path
from uuid import uuid4

from docx import Document

from app.database import Base, SessionLocal, engine
from app.models.chat import Artifact, Conversation, User
from app.services.chat_service import ChatService
from app.services.artifact_versioning import ArtifactVersionManager


def _seed_user_and_conversation(db):
    email = f'phase6_{uuid4().hex[:8]}@example.com'
    user = User(email=email, username='Phase6 User', hashed_password='demo-password')
    db.add(user)
    db.commit()
    db.refresh(user)

    conversation = Conversation(user_id=user.id, title='Phase 6 editing')
    db.add(conversation)
    db.commit()
    db.refresh(conversation)
    return user, conversation


def test_artifact_version_manager_creates_next_version(tmp_path):
    source = tmp_path / 'generated_proposal.docx'
    document = Document()
    document.add_heading('AI Strategy Proposal', level=0)
    document.add_paragraph('Original proposal text')
    document.save(str(source))

    manager = ArtifactVersionManager(base_dir=str(tmp_path))
    result = manager.apply_edit(str(source), 'Add a CFO summary and tighten the executive overview.')

    assert result['status'] == 'completed'
    assert result['version'] == 'v2'
    assert Path(result['file_path']).exists()
    assert result['file_path'] != str(source)
    assert 'v2' in Path(result['file_path']).name


def test_chat_service_updates_artifact_to_next_version():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        _, conversation = _seed_user_and_conversation(db)

        source_file = Path('storage/artifacts/generated_proposal.docx')
        source_file.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()
        doc.add_heading('AI Strategy Proposal', level=0)
        doc.add_paragraph('Baseline proposal')
        doc.save(str(source_file))

        artifact = Artifact(
            conversation_id=conversation.id,
            name='generated_proposal.docx',
            artifact_type='docx',
            version='v1',
            file_path=str(source_file),
            status='draft',
        )
        db.add(artifact)
        db.commit()
        db.refresh(artifact)

        service = ChatService(db)
        updated = service.apply_edit_to_artifact(artifact.id, 'Add a regional rollout breakdown and tighten the messaging.')

        assert updated.version == 'v2'
        assert updated.file_path != str(source_file)
        assert Path(updated.file_path).exists()
    finally:
        db.close()
