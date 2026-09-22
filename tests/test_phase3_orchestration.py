from pptx import Presentation

from app.agents.orchestrator import MultiAgentOrchestrator


def test_orchestrator_builds_response_for_document_request():
    orchestrator = MultiAgentOrchestrator()
    result = orchestrator.run(
        request='Research generative AI trends and create a proposal outline using the uploaded template.',
        conversation_id=1,
        uploaded_files=['sample.docx'],
        user_context={'theme': 'enterprise proposal'},
    )

    assert result['status'] == 'completed'
    assert 'supervisor' in result['route']
    assert 'research' in result['summary'].lower() or 'proposal' in result['summary'].lower()
    assert result['citations']


def test_orchestrator_summarizes_uploaded_ppt_main_content(tmp_path):
    ppt_path = tmp_path / 'project.pptx'
    prs = Presentation()

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = 'AI Strategy Overview'
    slide.placeholders[1].text = 'Vision\nRoadmap\nROI'

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = 'Business Impact'
    slide.shapes.placeholders[1].text = 'Lower cost\nFaster delivery\nHigher revenue'

    prs.save(ppt_path)

    orchestrator = MultiAgentOrchestrator()
    result = orchestrator.run(
        request='Research the ppt and give me main content',
        conversation_id=1,
        uploaded_files=[str(ppt_path)],
        user_context={'theme': 'enterprise proposal'},
    )

    summary = result['summary'].lower()
    assert 'main content' in summary or 'ai strategy overview' in summary or 'business impact' in summary
    assert 'slide' in summary or 'presentation' in summary
