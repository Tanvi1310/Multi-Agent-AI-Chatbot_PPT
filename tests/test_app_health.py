from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from app.document_processing.analysis_service import AnalysisService
from app.main import app


def test_root_health_and_api_health():
    client = TestClient(app)
    root = client.get('/')
    assert root.status_code == 200
    assert 'Multi-Agent AI Chatbot backend is running.' in root.json()['message']

    health = client.get('/health')
    assert health.status_code == 200
    assert health.json()['status'] == 'ok'

    api_health = client.get('/api/v1/health')
    assert api_health.status_code == 200
    assert api_health.json()['status'] == 'ok'


def test_document_analysis_extracts_style_and_text(tmp_path):
    doc_file = tmp_path / 'sample.docx'
    doc = Document()
    doc.add_heading('Executive Summary', 0)
    doc.add_paragraph('This proposal highlights growth in enterprise AI.')
    doc.save(doc_file)

    result = AnalysisService().analyze_file(str(doc_file))

    assert result['source_type'] == 'docx'
    assert result['title'] == 'Executive Summary'
    assert 'enterprise AI' in result['text']
    assert 'style' in result
    assert 'font_preferences' in result['style']
